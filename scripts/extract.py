#!/usr/bin/env python3
"""Extract every figure from the AEMO 2026 WEM ESOO data register.

For each figure sheet this writes ``<out>/<sheet-id>.json`` containing:
  * figure metadata (number, title, notes) read from the sheet,
  * the sheet's cell grid (values, with dates as ISO strings),
  * hidden rows/columns,
  * every native Excel chart on the sheet, parsed from its DrawingML XML:
    chart groups, series (with references resolved against the sheet and the
    chart's own cached values), resolved colours, line/marker styles, data
    labels, axes (titles, number formats, bounds, units, orientation), legend,
  * embedded pictures (for figures whose chart is a pasted image).

Also writes ``<out>/<sheet-id>.summary.txt`` (a compact human-readable digest)
and ``<out>/index.json`` (one entry per figure).

Usage: python3 -I scripts/extract.py <data-register.xlsx> <out-dir> [--pages figpages.json]
"""
import argparse
import colorsys
import datetime as dt
import json
import os
import posixpath
import re
import sys
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import column_index_from_string, get_column_letter

NS = {
    'm': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
    'rel': 'http://schemas.openxmlformats.org/package/2006/relationships',
    'xdr': 'http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'cs': 'http://schemas.microsoft.com/office/drawing/2012/chartStyle',
    'c15': 'http://schemas.microsoft.com/office/drawing/2012/chart',
    'c14': 'http://schemas.microsoft.com/office/drawing/2007/8/2/chart',
}
C = '{%s}' % NS['c']
A = '{%s}' % NS['a']
R = '{%s}' % NS['r']

SECTIONS = {
    'ES': 'Executive summary',
    'Ch 1': '1. Introduction',
    'Ch 2': '2. Consumption and demand forecasts',
    'Ch 3': '3. Supply forecasts',
    'Ch 4': '4. Electric Storage Resource and Demand Side Programme requirements',
    'Ch 5': '5. Peak and Flexible Reserve Capacity Target determination',
    'Ch 6': '6. Network congestion and sub-regional reliability shortfall assessment',
    'Ch 7': '7. Opportunities for investment in supply and transmission',
    'A1': 'A1. Consumption and demand forecasting inputs',
    'A2': 'A2. Assessment of supply projects',
    'A3': 'A3. Forecast supply capacity mix',
    'A4': 'A4. Reliability assessment framework, methodology and results',
}

PRST_COLORS = {'black': '000000', 'white': 'FFFFFF', 'red': 'FF0000', 'green': '008000', 'blue': '0000FF',
               'yellow': 'FFFF00', 'gray': '808080', 'grey': '808080', 'darkGray': 'A9A9A9', 'lightGray': 'D3D3D3'}


# ----------------------------------------------------------------------------------------------
# package helpers
class Package:
    def __init__(self, path):
        self.z = zipfile.ZipFile(path)
        self.names = set(self.z.namelist())

    def xml(self, name):
        return ET.fromstring(self.z.read(name))

    def rels(self, path):
        d, b = posixpath.dirname(path), posixpath.basename(path)
        rp = f'{d}/_rels/{b}.rels'
        if rp not in self.names:
            return {}
        out = {}
        for r in self.xml(rp).findall('rel:Relationship', NS):
            tgt = r.get('Target')
            if r.get('TargetMode') == 'External':
                full = tgt
            elif tgt.startswith('/'):
                full = tgt.lstrip('/')
            else:
                full = posixpath.normpath(posixpath.join(d, tgt))
            out[r.get('Id')] = (r.get('Type').rsplit('/', 1)[-1], full)
        return out


# ----------------------------------------------------------------------------------------------
# colours
def hex_to_rgb(h):
    h = h.strip('#')
    return tuple(int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))


def rgb_to_hex(rgb):
    return '#' + ''.join('%02X' % max(0, min(255, round(c * 255))) for c in rgb)


def _lin(c):
    return c / 12.92 if c <= 0.04045 else ((c + 0.055) / 1.055) ** 2.4


def _gam(c):
    return c * 12.92 if c <= 0.0031308 else 1.055 * (c ** (1 / 2.4)) - 0.055


class Theme:
    def __init__(self, root):
        self.scheme = {}
        cs = root.find('.//a:clrScheme', NS)
        if cs is not None:
            for el in cs:
                name = el.tag.split('}')[1]
                child = el[0]
                if child.tag == A + 'srgbClr':
                    self.scheme[name] = child.get('val')
                elif child.tag == A + 'sysClr':
                    self.scheme[name] = child.get('lastClr') or ('000000' if name == 'dk1' else 'FFFFFF')
        alias = {'tx1': 'dk1', 'bg1': 'lt1', 'tx2': 'dk2', 'bg2': 'lt2'}
        for k, v in alias.items():
            if v in self.scheme:
                self.scheme[k] = self.scheme[v]
        fonts = root.find('.//a:fontScheme', NS)
        self.minor_font = None
        if fonts is not None:
            mf = fonts.find('a:minorFont/a:latin', NS)
            self.minor_font = mf.get('typeface') if mf is not None else None


def apply_mods(hexval, mods):
    """Apply DrawingML colour transforms (lumMod/lumOff/tint/shade/satMod/alpha)."""
    r, g, b = hex_to_rgb(hexval)
    alpha = 1.0
    for tag, val in mods:
        v = int(val) / 100000
        if tag in ('lumMod', 'lumOff', 'satMod'):
            h, l, s = colorsys.rgb_to_hls(r, g, b)
            if tag == 'lumMod':
                l = l * v
            elif tag == 'lumOff':
                l = l + v
            else:
                s = s * v
            l = min(max(l, 0), 1)
            s = min(max(s, 0), 1)
            r, g, b = colorsys.hls_to_rgb(h, l, s)
        elif tag == 'tint':
            r, g, b = (_gam(_lin(c) * v + (1 - v)) for c in (r, g, b))
        elif tag == 'shade':
            r, g, b = (_gam(_lin(c) * v) for c in (r, g, b))
        elif tag == 'alpha':
            alpha = v
    return rgb_to_hex((r, g, b)), alpha


def resolve_color(el, theme):
    """el is a fill-ish element containing one colour child (srgbClr/schemeClr/...)."""
    if el is None:
        return None
    for child in el:
        tag = child.tag.split('}')[1]
        base = None
        if tag == 'srgbClr':
            base = child.get('val')
        elif tag == 'schemeClr':
            base = theme.scheme.get(child.get('val'))
            if base is None:
                return {'scheme': child.get('val')}
        elif tag == 'sysClr':
            base = child.get('lastClr') or '000000'
        elif tag == 'prstClr':
            base = PRST_COLORS.get(child.get('val'), '000000')
        elif tag == 'scrgbClr':
            base = rgb_to_hex(tuple(_gam(int(child.get(k)) / 100000) for k in ('r', 'g', 'b')))[1:]
        if base is None:
            continue
        mods = [(m.tag.split('}')[1], m.get('val')) for m in child if m.get('val') is not None]
        hx, alpha = apply_mods(base, mods)
        out = {'hex': hx}
        if alpha < 1:
            out['alpha'] = round(alpha, 3)
        if tag == 'schemeClr':
            out['scheme'] = child.get('val')
            if mods:
                out['mods'] = mods
        return out
    return None


def parse_fill(sppr, theme):
    """Return dict describing the fill of an spPr: {'none': True} | {'hex':..} | {'gradient': [...]} | None (auto)."""
    if sppr is None:
        return None
    for child in sppr:
        tag = child.tag.split('}')[1]
        if tag == 'noFill':
            return {'none': True}
        if tag == 'solidFill':
            return resolve_color(child, theme)
        if tag == 'gradFill':
            stops = []
            for gs in child.findall('a:gsLst/a:gs', NS):
                c = resolve_color(gs, theme)
                if c:
                    c['pos'] = int(gs.get('pos')) / 100000
                    stops.append(c)
            return {'gradient': stops, 'hex': stops[0]['hex'] if stops else None}
        if tag == 'pattFill':
            fg = resolve_color(child.find('a:fgClr', NS), theme)
            bg = resolve_color(child.find('a:bgClr', NS), theme)
            return {'pattern': child.get('prst'), 'hex': fg['hex'] if fg else None, 'bg': bg['hex'] if bg else None}
    return None


def parse_line(sppr, theme):
    if sppr is None:
        return None
    ln = sppr.find('a:ln', NS)
    if ln is None:
        return None
    out = {}
    if ln.get('w'):
        out['width_pt'] = round(int(ln.get('w')) / 12700, 2)
    for child in ln:
        tag = child.tag.split('}')[1]
        if tag == 'noFill':
            out['none'] = True
        elif tag == 'solidFill':
            col = resolve_color(child, theme)
            if col:
                out['color'] = col['hex']
                if 'alpha' in col:
                    out['alpha'] = col['alpha']
        elif tag == 'prstDash':
            out['dash'] = child.get('val')
        elif tag == 'custDash':
            out['dash'] = 'custom'
    return out or None


def rich_text(el):
    if el is None:
        return None
    paras = []
    for p in el.iter(A + 'p'):
        paras.append(''.join(t.text or '' for t in p.iter(A + 't')))
    txt = '\n'.join(x for x in paras if x is not None).strip()
    return txt or None


def font_size(el):
    """First defRPr/rPr sz under el (hundredths of a point) -> pt."""
    if el is None:
        return None
    for tag in ('defRPr', 'rPr'):
        for r in el.iter(A + tag):
            if r.get('sz'):
                return int(r.get('sz')) / 100
    return None


def text_color(el, theme):
    if el is None:
        return None
    for r in el.iter(A + 'defRPr'):
        sf = r.find('a:solidFill', NS)
        if sf is not None:
            c = resolve_color(sf, theme)
            return c['hex'] if c else None
    return None


def rotation(el):
    if el is None:
        return None
    bp = el.find('.//a:bodyPr', NS)
    if bp is not None and bp.get('rot') not in (None, '0', '-60000000'):
        return int(bp.get('rot')) / 60000
    return None


def val(el, path, cast=str):
    x = el.find(path, NS) if el is not None else None
    if x is None or x.get('val') is None:
        return None
    v = x.get('val')
    if cast is bool:
        return v in ('1', 'true')
    try:
        return cast(v)
    except ValueError:
        return v


# ----------------------------------------------------------------------------------------------
# workbook values
def excel_serial_to_iso(x, with_time=None):
    base = dt.datetime(1899, 12, 30)
    d = base + dt.timedelta(days=float(x))
    if with_time is None:
        with_time = abs(float(x) - round(float(x))) > 1e-9
    return d.strftime('%Y-%m-%dT%H:%M:%S') if with_time else d.strftime('%Y-%m-%d')


def is_date_format(code):
    if not code or code == 'General':
        return False
    c = re.sub(r'"[^"]*"|\\.|\[[^\]]*\]', '', code)
    return bool(re.search(r'[dmyhs]', c, re.I)) and not re.search(r'[0#]', c)


def norm_value(v):
    if isinstance(v, dt.datetime):
        if v.hour == v.minute == v.second == 0 and v.microsecond == 0:
            return v.strftime('%Y-%m-%d')
        return v.strftime('%Y-%m-%dT%H:%M:%S')
    if isinstance(v, dt.date):
        return v.strftime('%Y-%m-%d')
    if isinstance(v, dt.time):
        return v.strftime('%H:%M:%S')
    if isinstance(v, dt.timedelta):
        return v.total_seconds() / 86400
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
        return int(v)
    if isinstance(v, str):
        return v
    return v


class SheetGrid:
    """All cell values of one sheet, row-major from A1. Rows/cols are 1-based in API."""

    def __init__(self, ws, fmt_rows=400):
        self.rows = []
        self.formats = {}
        for i, row in enumerate(ws.iter_rows(), start=1):
            vals = []
            for j, cell in enumerate(row, start=1):
                v = getattr(cell, 'value', None)
                vals.append(norm_value(v))
                if i <= fmt_rows and v is not None:
                    nf = getattr(cell, 'number_format', None)
                    if nf and nf != 'General':
                        self.formats[f'{get_column_letter(j)}{i}'] = nf
            while vals and vals[-1] is None:
                vals.pop()
            self.rows.append(vals)
        while self.rows and not self.rows[-1]:
            self.rows.pop()

    def get(self, r, c):
        if r - 1 < len(self.rows):
            row = self.rows[r - 1]
            if c - 1 < len(row):
                return row[c - 1]
        return None


def sheet_hidden(pkg, path):
    """Hidden rows and columns of a sheet (1-based indices), parsed from the sheet XML."""
    hidden_rows, hidden_cols = [], []
    with pkg.z.open(path) as f:
        for ev, el in ET.iterparse(f, events=('end',)):
            tag = el.tag.split('}')[1]
            if tag == 'row':
                if el.get('hidden') in ('1', 'true'):
                    hidden_rows.append(int(el.get('r')))
                el.clear()
            elif tag == 'col':
                if el.get('hidden') in ('1', 'true'):
                    hidden_cols.extend(range(int(el.get('min')), int(el.get('max')) + 1))
            elif tag == 'sheetData':
                el.clear()
    return hidden_rows, hidden_cols


REF_RE = re.compile(r"^(?:\[(?P<ext>\d+)\])?(?:'(?P<q>(?:[^']|'')+)'|(?P<s>[^!']+))!(?P<rng>.+)$")


def parse_ref(f):
    """'Sheet'!$B$2:$C$9 -> (sheet, (r1,c1,r2,c2)) ; returns error string for #REF!/external."""
    if f is None:
        return None, None, None
    f = f.strip()
    if f.startswith('(') and f.endswith(')'):
        parts = [p.strip() for p in f[1:-1].split(',')]
        res = [parse_ref(p) for p in parts]
        return 'union', res, None
    m = REF_RE.match(f)
    if not m:
        return None, None, 'unparsed'
    if m.group('ext'):
        return None, None, 'external'
    sheet = (m.group('q') or m.group('s')).replace("''", "'")
    rng = m.group('rng')
    if '#REF' in rng:
        return sheet, None, '#REF!'
    cells = rng.replace('$', '').split(':')
    def rc(x):
        mm = re.match(r'([A-Z]+)(\d+)', x)
        if not mm:
            return None
        return int(mm.group(2)), column_index_from_string(mm.group(1))
    a = rc(cells[0])
    b = rc(cells[1]) if len(cells) > 1 else a
    if a is None or b is None:
        return sheet, None, 'unparsed'
    return sheet, (min(a[0], b[0]), min(a[1], b[1]), max(a[0], b[0]), max(a[1], b[1])), None


# ----------------------------------------------------------------------------------------------
# charts
def parse_cache(ref_el):
    """Parse numCache/strCache/multiLvlStrCache under a *Ref or *Lit element."""
    out = {}
    if ref_el is None:
        return out
    f = ref_el.find('c:f', NS)
    if f is not None:
        out['ref'] = f.text
    for cache_tag in ('numCache', 'strCache', 'numLit', 'strLit'):
        cache = ref_el.find(f'c:{cache_tag}', NS) if not ref_el.tag.endswith(('numLit', 'strLit')) else ref_el
        if cache is None:
            continue
        n = val(cache, 'c:ptCount', int) or 0
        pts = [None] * n
        for pt in cache.findall('c:pt', NS):
            i = int(pt.get('idx'))
            v = pt.find('c:v', NS)
            if i < n and v is not None:
                pts[i] = v.text
        fc = cache.find('c:formatCode', NS)
        code = fc.text if fc is not None else None
        if cache_tag in ('numCache', 'numLit'):
            conv = []
            for p in pts:
                try:
                    conv.append(float(p) if p is not None else None)
                except ValueError:
                    conv.append(p)
            pts = conv
            if code and is_date_format(code):
                out['cacheDates'] = [excel_serial_to_iso(p) if isinstance(p, float) else p for p in pts]
        out['cache'] = [norm_value(p) for p in pts]
        if code:
            out['formatCode'] = code
        break
    ml = ref_el.find('c:multiLvlStrCache', NS)
    if ml is not None:
        n = val(ml, 'c:ptCount', int) or 0
        levels = []
        for lvl in ml.findall('c:lvl', NS):
            pts = [None] * n
            for pt in lvl.findall('c:pt', NS):
                i = int(pt.get('idx'))
                v = pt.find('c:v', NS)
                if i < n and v is not None:
                    pts[i] = v.text
            levels.append(pts)
        out['cacheLevels'] = levels  # level 0 is the innermost (closest to axis)
        out['multiLevel'] = True
    return out


def parse_data_ref(el):
    """c:cat / c:val / c:xVal / c:yVal / c:bubbleSize / c:tx."""
    if el is None:
        return None
    for tag in ('numRef', 'strRef', 'multiLvlStrRef', 'numLit', 'strLit'):
        x = el.find(f'c:{tag}', NS)
        if x is not None:
            d = parse_cache(x)
            d['kind'] = tag
            return d
    v = el.find('c:v', NS)
    if v is not None:
        return {'kind': 'literal', 'cache': [v.text]}
    return None


def parse_dlbls(dl, theme):
    if dl is None:
        return None
    out = {}
    for k in ('showVal', 'showSerName', 'showCatName', 'showPercent', 'showLegendKey', 'showBubbleSize'):
        v = val(dl, f'c:{k}', bool)
        if v is not None:
            out[k] = v
    if val(dl, 'c:delete', bool):
        out['delete'] = True
    pos = val(dl, 'c:dLblPos')
    if pos:
        out['pos'] = pos
    nf = dl.find('c:numFmt', NS)
    if nf is not None:
        out['numFmt'] = nf.get('formatCode')
    fs = font_size(dl.find('c:txPr', NS))
    if fs:
        out['fontSize'] = fs
    tc = text_color(dl.find('c:txPr', NS), theme)
    if tc:
        out['color'] = tc
    pts = []
    for d in dl.findall('c:dLbl', NS):
        p = {'idx': val(d, 'c:idx', int)}
        if val(d, 'c:delete', bool):
            p['delete'] = True
        for k in ('showVal', 'showSerName', 'showCatName'):
            v = val(d, f'c:{k}', bool)
            if v is not None:
                p[k] = v
        t = d.find('c:tx', NS)
        if t is not None:
            p['text'] = rich_text(t)
        if val(d, 'c:dLblPos'):
            p['pos'] = val(d, 'c:dLblPos')
        pts.append(p)
    if pts:
        out['points'] = pts
    return out or None


def parse_marker(m, theme):
    if m is None:
        return None
    out = {}
    s = val(m, 'c:symbol')
    if s:
        out['symbol'] = s
    sz = val(m, 'c:size', int)
    if sz:
        out['size'] = sz
    sp = m.find('c:spPr', NS)
    f = parse_fill(sp, theme)
    if f:
        out['fill'] = f
    ln = parse_line(sp, theme)
    if ln:
        out['line'] = ln
    return out or None


def parse_series(ser, group_type, theme):
    s = {'idx': val(ser, 'c:idx', int), 'order': val(ser, 'c:order', int), 'type': group_type}
    tx = ser.find('c:tx', NS)
    if tx is not None:
        s['name'] = parse_data_ref(tx)
    sp = ser.find('c:spPr', NS)
    fill = parse_fill(sp, theme)
    line = parse_line(sp, theme)
    if fill is not None:
        s['fill'] = fill
    if line is not None:
        s['line'] = line
    mk = parse_marker(ser.find('c:marker', NS), theme)
    if mk:
        s['marker'] = mk
    for k in ('smooth', 'invertIfNegative'):
        v = val(ser, f'c:{k}', bool)
        if v is not None:
            s[k] = v
    if val(ser, 'c:explosion', int):
        s['explosion'] = val(ser, 'c:explosion', int)
    for tag, key in (('c:cat', 'cat'), ('c:val', 'val'), ('c:xVal', 'x'), ('c:yVal', 'y'), ('c:bubbleSize', 'size')):
        d = parse_data_ref(ser.find(tag, NS))
        if d is not None:
            s[key] = d
    dps = []
    for dp in ser.findall('c:dPt', NS):
        d = {'idx': val(dp, 'c:idx', int)}
        f = parse_fill(dp.find('c:spPr', NS), theme)
        ln = parse_line(dp.find('c:spPr', NS), theme)
        if f:
            d['fill'] = f
        if ln:
            d['line'] = ln
        mk = parse_marker(dp.find('c:marker', NS), theme)
        if mk:
            d['marker'] = mk
        dps.append(d)
    if dps:
        s['points'] = dps
    dl = parse_dlbls(ser.find('c:dLbls', NS), theme)
    if dl:
        s['dataLabels'] = dl
    tls = []
    for tl in ser.findall('c:trendline', NS):
        tls.append({'type': val(tl, 'c:trendlineType'), 'line': parse_line(tl.find('c:spPr', NS), theme),
                    'name': (tl.find('c:name', NS).text if tl.find('c:name', NS) is not None else None)})
    if tls:
        s['trendlines'] = tls
    ebs = []
    for eb in ser.findall('c:errBars', NS):
        e = {'dir': val(eb, 'c:errDir'), 'barType': val(eb, 'c:errBarType'), 'valType': val(eb, 'c:errValType'),
             'noEndCap': val(eb, 'c:noEndCap', bool), 'value': val(eb, 'c:val', float),
             'line': parse_line(eb.find('c:spPr', NS), theme)}
        for k in ('plus', 'minus'):
            d = parse_data_ref(eb.find(f'c:{k}', NS))
            if d:
                e[k] = d
        ebs.append(e)
    if ebs:
        s['errorBars'] = ebs
    return s


def parse_axis(ax, theme):
    kind = ax.tag.split('}')[1]
    a = {'id': val(ax, 'c:axId'), 'kind': kind, 'pos': val(ax, 'c:axPos'), 'delete': bool(val(ax, 'c:delete', bool)),
         'crossAx': val(ax, 'c:crossAx')}
    sc = ax.find('c:scaling', NS)
    a['orientation'] = val(sc, 'c:orientation')
    for k in ('min', 'max', 'logBase'):
        v = val(sc, f'c:{k}', float)
        if v is not None:
            a[k] = v
    for k in ('majorUnit', 'minorUnit'):
        v = val(ax, f'c:{k}', float)
        if v is not None:
            a[k] = v
    for k in ('crosses', 'crossBetween', 'tickLblPos', 'majorTickMark', 'minorTickMark', 'baseTimeUnit',
              'majorTimeUnit', 'lblAlgn'):
        v = val(ax, f'c:{k}')
        if v is not None:
            a[k] = v
    for k in ('crossesAt', ):
        v = val(ax, f'c:{k}', float)
        if v is not None:
            a[k] = v
    for k in ('tickLblSkip', 'tickMarkSkip', 'lblOffset'):
        v = val(ax, f'c:{k}', int)
        if v is not None:
            a[k] = v
    du = ax.find('c:dispUnits', NS)
    if du is not None:
        b = du.find('c:builtInUnit', NS)
        cu = du.find('c:custUnit', NS)
        units = {'hundreds': 1e2, 'thousands': 1e3, 'tenThousands': 1e4, 'hundredThousands': 1e5, 'millions': 1e6,
                 'tenMillions': 1e7, 'hundredMillions': 1e8, 'billions': 1e9, 'trillions': 1e12}
        if b is not None:
            a['dispUnits'] = b.get('val')
            a['dispUnitsDivisor'] = units.get(b.get('val'))
        elif cu is not None:
            a['dispUnitsDivisor'] = float(cu.get('val'))
        lbl = du.find('c:dispUnitsLbl', NS)
        if lbl is not None:
            a['dispUnitsLabel'] = rich_text(lbl.find('c:tx', NS)) or a.get('dispUnits')
    nf = ax.find('c:numFmt', NS)
    if nf is not None:
        a['numFmt'] = nf.get('formatCode')
        a['numFmtSourceLinked'] = nf.get('sourceLinked') == '1'
    t = ax.find('c:title', NS)
    if t is not None:
        a['title'] = rich_text(t.find('c:tx', NS)) or (t.find('.//c:f', NS).text if t.find('.//c:f', NS) is not None else None)
        a['titleRotation'] = rotation(t.find('c:tx/c:rich', NS))
    a['majorGridlines'] = ax.find('c:majorGridlines', NS) is not None
    if a['majorGridlines']:
        gl = parse_line(ax.find('c:majorGridlines/c:spPr', NS), theme)
        if gl:
            a['gridline'] = gl
    a['minorGridlines'] = ax.find('c:minorGridlines', NS) is not None
    a['line'] = parse_line(ax.find('c:spPr', NS), theme)
    tx = ax.find('c:txPr', NS)
    a['labelFontSize'] = font_size(tx)
    a['labelRotation'] = rotation(tx)
    a['labelColor'] = text_color(tx, theme)
    return {k: v for k, v in a.items() if v is not None}


def load_color_style(pkg, path, theme):
    """Chart colour style (colorsN.xml): list of base colours and variations for automatic fills."""
    if path is None or path not in pkg.names:
        return None
    root = pkg.xml(path)
    bases = []
    for el in root:
        tag = el.tag.split('}')[1]
        if tag in ('schemeClr', 'srgbClr', 'sysClr'):
            wrapper = ET.Element('w')
            wrapper.append(el)
            c = resolve_color(wrapper, theme)
            if c:
                bases.append((el, c))
    variations = []
    for v in root.findall('cs:variation', NS):
        variations.append([(m.tag.split('}')[1], m.get('val')) for m in v])
    return {'meth': root.get('meth'), 'bases': bases, 'variations': variations or [[]]}


def auto_color(style, theme, idx):
    if not style or not style['bases']:
        accents = [theme.scheme.get(f'accent{i}') for i in range(1, 7)]
        base = accents[idx % 6]
        return '#' + base if base else None
    bases, variations = style['bases'], style['variations']
    el, _ = bases[idx % len(bases)]
    var = variations[(idx // len(bases)) % len(variations)]
    tag = el.tag.split('}')[1]
    base = theme.scheme.get(el.get('val')) if tag == 'schemeClr' else (el.get('val') if tag == 'srgbClr' else el.get('lastClr'))
    mods = [(m.tag.split('}')[1], m.get('val')) for m in el] + var
    return apply_mods(base, mods)[0]


def parse_chart(pkg, path, theme):
    root = pkg.xml(path)
    chart = root.find('c:chart', NS)
    crels = pkg.rels(path)
    color_style = None
    chart_theme = theme
    for typ, tgt in crels.values():
        if typ == 'chartColorStyle':
            color_style = load_color_style(pkg, tgt, theme)
        if typ == 'themeOverride' and tgt in pkg.names:
            chart_theme = Theme(pkg.xml(tgt))
            for k, v in theme.scheme.items():
                chart_theme.scheme.setdefault(k, v)
    th = chart_theme
    spec = {'file': posixpath.basename(path)}
    # shapes drawn inside the chart area (text boxes, arrows, braces): positions are fractions of the chart
    shapes = []
    for typ, tgt in crels.values():
        if typ != 'chartUserShapes' or tgt not in pkg.names:
            continue
        CDR = '{http://schemas.openxmlformats.org/drawingml/2006/chartDrawing}'
        for anc in pkg.xml(tgt):
            fr, to = anc.find(CDR + 'from'), anc.find(CDR + 'to')
            xy = lambda e: [round(float(e.find(CDR + 'x').text), 4), round(float(e.find(CDR + 'y').text), 4)] if e is not None else None
            for shp in anc.iter():
                tag = shp.tag.split('}')[1] if '}' in shp.tag else shp.tag
                if tag not in ('sp', 'cxnSp', 'pic'):
                    continue
                geom = shp.find('.//a:prstGeom', NS)
                txt = rich_text(shp.find(CDR + 'txBody')) if shp.find(CDR + 'txBody') is not None else None
                shapes.append({'kind': tag, 'geom': geom.get('prst') if geom is not None else None, 'text': txt,
                               'from': xy(fr), 'to': xy(to),
                               'fill': parse_fill(shp.find(CDR + 'spPr'), th), 'line': parse_line(shp.find(CDR + 'spPr'), th)})
    if shapes:
        spec['userShapes'] = shapes
    t = chart.find('c:title', NS)
    if t is not None:
        spec['title'] = rich_text(t.find('c:tx', NS)) or (t.find('.//c:f', NS).text if t.find('.//c:f', NS) is not None else '(auto title)')
    spec['autoTitleDeleted'] = val(chart, 'c:autoTitleDeleted', bool)
    spec['plotVisOnly'] = val(chart, 'c:plotVisOnly', bool)
    spec['dispBlanksAs'] = val(chart, 'c:dispBlanksAs')
    spec['fontSize'] = font_size(root.find('c:txPr', NS))
    lat = root.find('c:txPr//a:latin', NS)
    spec['font'] = lat.get('typeface') if lat is not None else th.minor_font
    lg = chart.find('c:legend', NS)
    if lg is not None:
        L = {'pos': val(lg, 'c:legendPos'), 'overlay': val(lg, 'c:overlay', bool), 'fontSize': font_size(lg.find('c:txPr', NS))}
        dele = []
        for le in lg.findall('c:legendEntry', NS):
            if val(le, 'c:delete', bool):
                dele.append(val(le, 'c:idx', int))
        if dele:
            L['deletedEntries'] = dele  # indices into the legend's entry list (series idx order)
        spec['legend'] = {k: v for k, v in L.items() if v is not None}
    else:
        spec['legend'] = None
    plot = chart.find('c:plotArea', NS)
    groups, axes, filtered = [], [], []
    for child in plot:
        tag = child.tag.split('}')[1]
        if tag.endswith('Chart'):
            gtype = tag[:-5]  # bar, line, area, scatter, pie, doughnut, bubble, radar, stock, bar3D...
            g = {'type': gtype}
            for k in ('barDir', 'grouping', 'scatterStyle', 'radarStyle', 'shape'):
                v = val(child, f'c:{k}')
                if v is not None:
                    g[k] = v
            for k in ('gapWidth', 'overlap', 'firstSliceAng', 'holeSize', 'bubbleScale'):
                v = val(child, f'c:{k}', int)
                if v is not None:
                    g[k] = v
            for k in ('varyColors', ):
                v = val(child, f'c:{k}', bool)
                if v is not None:
                    g[k] = v
            if child.find('c:dropLines', NS) is not None:
                g['dropLines'] = True
            if child.find('c:hiLowLines', NS) is not None:
                g['hiLowLines'] = True
            if child.find('c:upDownBars', NS) is not None:
                g['upDownBars'] = True
            g['axIds'] = [a.get('val') for a in child.findall('c:axId', NS)]
            gd = parse_dlbls(child.find('c:dLbls', NS), th)
            if gd and any(gd.get(k) for k in ('showVal', 'showSerName', 'showCatName', 'showPercent')):
                g['dataLabels'] = gd
            g['series'] = [parse_series(s, gtype, th) for s in child.findall('c:ser', NS)]
            for s in g['series']:
                if s.get('fill') is None and gtype not in ('line', 'scatter'):
                    s['fill'] = {'hex': auto_color(color_style, th, s['idx'] or 0), 'auto': True}
                if gtype in ('line', 'scatter', 'radar') and not (s.get('line') or {}).get('color') and not (s.get('line') or {}).get('none'):
                    s.setdefault('line', {})
                    s['line']['color'] = auto_color(color_style, th, s['idx'] or 0)
                    s['line']['auto'] = True
            # filtered (hidden via chart filter) series live in extLst
            for ext in child.iter():
                et = ext.tag.split('}')[1] if '}' in ext.tag else ext.tag
                if et.startswith('filtered') and et.endswith('Series'):
                    for fs in ext.findall('.//c:ser', NS):
                        fsd = parse_series(fs, gtype, th)
                        filtered.append({'group': gtype, 'idx': fsd.get('idx'), 'name': fsd.get('name'),
                                         'val': (fsd.get('val') or fsd.get('y') or {}).get('ref')})
            groups.append(g)
        elif tag in ('catAx', 'valAx', 'dateAx', 'serAx'):
            axes.append(parse_axis(child, th))
        elif tag == 'dTable':
            spec['dataTable'] = True
    # top-level filtered series (c15 ext at plot area level)
    for ext in plot.iter():
        et = ext.tag.split('}')[1] if '}' in ext.tag else ext.tag
        if et.startswith('filtered') and et.endswith('Series') and not any(f for f in filtered):
            pass
    spec['groups'] = groups
    spec['axes'] = axes
    if filtered:
        spec['filteredSeries'] = filtered
    # which value axes are secondary (the second set of axes, usually crosses="max")
    ax_by_id = {a['id']: a for a in axes}
    if groups:
        # Excel writes the primary category/value axis pair first; any group not on it is secondary
        primary = {a['id'] for a in axes[:2]} if len(axes) > 2 else set(groups[0]['axIds'])
        for g in groups:
            g['secondary'] = not (set(g['axIds']) & primary)
            g['axes'] = {ax_by_id[i]['kind'] + ('' if ax_by_id[i]['kind'] != 'valAx' else '') + ':' + i: ax_by_id[i].get('pos') for i in g['axIds'] if i in ax_by_id}
    return spec


# ----------------------------------------------------------------------------------------------
# resolve chart references against sheet grids
def resolve_ref(d, grids, hidden, default_sheet):
    """Resolve a parsed data ref dict in place: adds 'values' (and 'hiddenIdx')."""
    if not d or 'ref' not in d:
        return
    sheet, rng, err = parse_ref(d['ref'])
    if sheet == 'union':
        vals, hid = [], []
        for (sh, rg, er) in rng:
            if rg is None or sh not in grids:
                d['refError'] = er or 'unresolved'
                return
            v, h = _cells(grids[sh], rg, hidden.get(sh, ([], [])), len(vals))
            vals += v
            hid += h
        d['values'] = vals
        if hid:
            d['hiddenIdx'] = hid
        return
    if err:
        d['refError'] = err
        return
    if sheet not in grids:
        d['refError'] = 'sheet not loaded'
        return
    if d.get('kind') == 'multiLvlStrRef':
        r1, c1, r2, c2 = rng
        g = grids[sheet]
        if (c2 - c1) >= (r2 - r1):  # levels are rows
            d['levels'] = [[g.get(r, c) for c in range(c1, c2 + 1)] for r in range(r1, r2 + 1)]
        else:
            d['levels'] = [[g.get(r, c) for r in range(r1, r2 + 1)] for c in range(c1, c2 + 1)]
        hr, hc = hidden.get(sheet, ([], []))
        hid = [i for i, c in enumerate(range(c1, c2 + 1)) if c in hc] if (c2 - c1) >= (r2 - r1) else [i for i, r in enumerate(range(r1, r2 + 1)) if r in hr]
        if hid:
            d['hiddenIdx'] = hid
        return
    v, h = _cells(grids[sheet], rng, hidden.get(sheet, ([], [])), 0)
    d['values'] = v
    if h:
        d['hiddenIdx'] = h


def _cells(grid, rng, hidden, offset):
    r1, c1, r2, c2 = rng
    hr, hc = set(hidden[0]), set(hidden[1])
    vals, hid = [], []
    if r1 == r2 or c1 == c2:
        coords = [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    else:  # 2-D range: row-major
        coords = [(r, c) for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
    for i, (r, c) in enumerate(coords):
        v = grid.get(r, c)
        if isinstance(v, str) and v.startswith('#') and v.endswith(('!', 'A', '?')):
            v = None
        vals.append(v)
        if r in hr or c in hc:
            hid.append(offset + i)
    return vals, hid


# ----------------------------------------------------------------------------------------------
def figure_meta(sheet, grid):
    title_cell = grid.get(1, 1)
    full = str(title_cell).strip() if title_cell else sheet
    m = re.match(r'^Figure\s+([A-Z]?\d+[a-z]?)\s*[-–—:]\s*(.*)$', full, re.S)
    number, title = (m.group(1), m.group(2).strip()) if m else (None, full)
    if sheet.startswith('ES'):
        sec = 'ES'
    elif sheet.startswith('Ch '):
        sec = sheet.split('_')[0]
    else:
        sec = sheet.split('_')[0]
    doc = 'appendices' if re.match(r'^A\d', sheet) else 'report'
    notes = []
    for i, row in enumerate(grid.rows, start=1):
        if i <= 2:
            continue
        strs = [v.strip() for v in row if isinstance(v, str) and v.strip()]
        if not strs:
            continue
        has_num = any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in row)
        for s in strs:
            heading = re.match(r'^(Notes?|Sources?)\b', s, re.I)
            # a note is a text-only row's bullet, sentence or Note/Source line, or any very long text cell
            sentence = (s.endswith('.') or s.endswith('.)')) and len(s.split()) >= 4
            bullet = re.match(r'^[–\-▪•·*]\s*\S', s) and len(s.split()) >= 3
            if (heading or (not has_num and len(strs) <= 2 and (sentence or bullet)) or len(s) > 90) \
                    and not re.match(r'^[\d.,%-]+$', s) and s not in notes:
                notes.append(s)
    return {'number': number, 'title': title, 'fullTitle': full, 'section': sec,
            'sectionName': SECTIONS.get(sec, sec), 'doc': doc, 'notes': notes}


def sheet_id(name):
    return name.replace(' ', '')  # 'Ch 2_F.7' -> 'Ch2_F.7'


def fmt_num(v):
    if isinstance(v, float):
        return f'{v:.6g}'
    return str(v)


def summarize(fig):
    L = []
    L.append(f"{fig['id']}  (sheet '{fig['sheet']}')  kind={fig['kind']}")
    L.append(f"Figure {fig['number']} [{fig['doc']}]: {fig['title']}")
    L.append(f"Section: {fig['sectionName']}")
    if fig.get('pdf'):
        L.append(f"Reference page image(s): {', '.join(fig['pdf']['images'])}")
    for p in fig.get('pictures', []):
        L.append(f"Embedded picture: {p['file']}  anchor {p['from']}->{p['to']}")
    if fig['hiddenRows'] or fig['hiddenCols']:
        L.append(f"Hidden rows: {fig['hiddenRows'][:40]}  hidden cols: {[get_column_letter(c) for c in fig['hiddenCols']]}")
    for n in fig['notes']:
        L.append(f"NOTE: {n}")
    for ci, ch in enumerate(fig['charts']):
        L.append('')
        L.append(f"=== Chart {ci + 1}: {ch['file']}  anchor from {ch['anchor']['from']} to {ch['anchor']['to']}"
                 + (' HIDDEN-SHAPE' if ch['anchor'].get('hidden') else '')
                 + (' ZERO-SIZE: invisible leftover object in the workbook - do NOT recreate' if ch['anchor'].get('zeroSize') else ''))
        L.append(f"title={ch.get('title')!r} plotVisOnly={ch.get('plotVisOnly')} dispBlanksAs={ch.get('dispBlanksAs')} "
                 f"font={ch.get('font')} fontSize={ch.get('fontSize')} legend={ch.get('legend')}")
        for ax in ch['axes']:
            keys = ['kind', 'id', 'pos', 'delete', 'orientation', 'min', 'max', 'majorUnit', 'numFmt', 'dispUnits', 'dispUnitsDivisor', 'title',
                    'crosses', 'crossesAt', 'crossBetween', 'tickLblPos', 'majorGridlines', 'labelRotation',
                    'labelFontSize', 'baseTimeUnit', 'majorTimeUnit', 'tickLblSkip', 'logBase']
            L.append('  AXIS ' + ' '.join(f"{k}={ax[k]!r}" for k in keys if k in ax))
        for g in ch['groups']:
            gk = {k: v for k, v in g.items() if k not in ('series', 'axes')}
            L.append(f"  GROUP {gk}")
            for s in g['series']:
                nm = s.get('name') or {}
                name = (nm.get('values') or nm.get('cache') or [None])[0]
                fill = s.get('fill') or {}
                line = s.get('line') or {}
                sty = []
                if fill:
                    sty.append('fill=' + ('none' if fill.get('none') else str(fill.get('hex'))) + (' (auto)' if fill.get('auto') else '') + (f" alpha={fill['alpha']}" if fill.get('alpha') else ''))
                if line:
                    sty.append('line=' + ('none' if line.get('none') else f"{line.get('color')} w={line.get('width_pt')}pt dash={line.get('dash')}") + (' (auto)' if line.get('auto') else '') + (f" alpha={line['alpha']}" if line.get('alpha') else ''))
                if s.get('marker'):
                    sty.append(f"marker={s['marker']}")
                if s.get('smooth'):
                    sty.append('smooth')
                if s.get('dataLabels'):
                    sty.append(f"dLbls={s['dataLabels']}")
                if s.get('points'):
                    sty.append(f"dPt overrides={len(s['points'])}")
                if s.get('trendlines'):
                    sty.append(f"trendlines={s['trendlines']}")
                if s.get('errorBars'):
                    sty.append('errorBars')
                L.append(f"    SERIES idx={s['idx']} order={s['order']} name={name!r} ({nm.get('ref')}) {' '.join(sty)}")
                for key in ('cat', 'val', 'x', 'y', 'size'):
                    d = s.get(key)
                    if not d:
                        continue
                    vals = d.get('values') if d.get('values') is not None else d.get('levels') if d.get('levels') is not None else d.get('cache')
                    src = 'sheet' if (d.get('values') is not None or d.get('levels') is not None) else 'CACHE'
                    if d.get('refError'):
                        src = f"CACHE (ref error {d['refError']})"
                    n = len(vals) if vals is not None else 0
                    if vals and isinstance(vals[0], list):
                        prev = ' / '.join(str([fmt_num(x) for x in lv[:6]]) for lv in vals)
                    else:
                        prev = ', '.join(fmt_num(x) for x in (vals or [])[:8]) + (' …' if n > 8 else '')
                    extra = f" fmt={d['formatCode']!r}" if d.get('formatCode') else ''
                    hid = f" hiddenIdx={d['hiddenIdx'][:20]}" if d.get('hiddenIdx') else ''
                    mism = ''
                    if d.get('values') is not None and d.get('cache') is not None and d.get('cacheMismatch'):
                        mism = f" !! cache differs from sheet at {d['cacheMismatch']} points"
                    L.append(f"      {key}: {d.get('ref')} [{src}, n={n}]{extra}{hid}{mism}: {prev}")
        if ch.get('filteredSeries'):
            L.append(f"  FILTERED (hidden by chart filter, not displayed): {ch['filteredSeries']}")
        for sh in ch.get('userShapes', []):
            L.append(f"  DRAWN SHAPE {sh['kind']}/{sh['geom']} from {sh['from']} to {sh['to']} (fractions of chart)"
                     + (f" text={sh['text']!r}" if sh.get('text') else ''))
    L.append('')
    L.append('=== Sheet cells (first 60 non-empty rows, row: col=value) ===')
    shown = 0
    for i, row in enumerate(fig['grid'], start=1):
        if i <= 2 or not any(v is not None for v in row):
            continue
        cells = [f"{get_column_letter(j + 1)}={fmt_num(v)[:48]}" for j, v in enumerate(row) if v is not None]
        L.append(f"{i}: " + ' | '.join(cells)[:600])
        shown += 1
        if shown >= 60:
            L.append(f"... ({len(fig['grid'])} rows in total)")
            break
    return '\n'.join(L)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('xlsx')
    ap.add_argument('out')
    ap.add_argument('--pages', help='figpages.json mapping figure number -> PDF pages')
    ap.add_argument('--page-dir', default='work/ref/pages')
    args = ap.parse_args()
    os.makedirs(args.out, exist_ok=True)
    pages = json.load(open(args.pages)) if args.pages else {}

    pkg = Package(args.xlsx)
    theme = Theme(pkg.xml('xl/theme/theme1.xml'))
    wbx = pkg.xml('xl/workbook.xml')
    wbrels = pkg.rels('xl/workbook.xml')
    sheets = [(s.get('name'), wbrels[s.get('{%s}id' % NS['r'])][1]) for s in wbx.find('m:sheets', NS)]

    print('loading values…', file=sys.stderr)
    wb = openpyxl.load_workbook(args.xlsx, read_only=True, data_only=True)
    grids, hidden = {}, {}
    for name, path in sheets:
        grids[name] = SheetGrid(wb[name])
        hidden[name] = sheet_hidden(pkg, path)
    print('parsing charts…', file=sys.stderr)

    index = []
    skip = {'Version Control', 'Contents', 'Appendices'}
    for name, path in sheets:
        if name in skip:
            continue
        grid = grids[name]
        fig = {'id': sheet_id(name), 'sheet': name}
        fig.update(figure_meta(name, grid))
        fig['hiddenRows'], fig['hiddenCols'] = hidden[name]
        fig['charts'], fig['pictures'] = [], []
        for typ, tgt in pkg.rels(path).values():
            if typ != 'drawing':
                continue
            drels = pkg.rels(tgt)
            for anchor in pkg.xml(tgt):
                frm, to = anchor.find('xdr:from', NS), anchor.find('xdr:to', NS)
                pos = lambda e: [int(e.find('xdr:col', NS).text), int(e.find('xdr:row', NS).text)] if e is not None else None
                for gf in anchor.iter('{%s}graphicFrame' % NS['xdr']):
                    ch = gf.find('.//c:chart', NS)
                    if ch is None:
                        continue
                    nv = gf.find('xdr:nvGraphicFramePr/xdr:cNvPr', NS)
                    cpath = drels[ch.get(R + 'id')][1]
                    spec = parse_chart(pkg, cpath, theme)
                    spec['anchor'] = {'from': pos(frm), 'to': pos(to), 'name': nv.get('name') if nv is not None else None}
                    if nv is not None and nv.get('hidden') in ('1', 'true'):
                        spec['anchor']['hidden'] = True
                    if frm is not None and to is not None:
                        off = lambda e, k: int(e.find(f'xdr:{k}', NS).text)
                        h = (off(to, 'row') - off(frm, 'row')) * 15 + (off(to, 'rowOff') - off(frm, 'rowOff')) / 12700
                        w = (off(to, 'col') - off(frm, 'col')) * 48 + (off(to, 'colOff') - off(frm, 'colOff')) / 12700
                        if h <= 2 or w <= 2:
                            spec['anchor']['zeroSize'] = True  # invisible leftover object in the workbook
                    for g in spec['groups']:
                        for s in g['series']:
                            for key in ('name', 'cat', 'val', 'x', 'y', 'size'):
                                d = s.get(key)
                                if d:
                                    resolve_ref(d, grids, hidden, name)
                                    if d.get('values') is not None and d.get('cache') is not None and key != 'name':
                                        mism = 0
                                        vis = d['values']
                                        if d.get('hiddenIdx') and len(d['cache']) == len(vis) - len(d['hiddenIdx']):
                                            hid = set(d['hiddenIdx'])  # Excel caches only visible cells
                                            vis = [v for i, v in enumerate(vis) if i not in hid]
                                        for a, b in zip(vis, d['cache']):
                                            if isinstance(a, (int, float)) and isinstance(b, (int, float)):
                                                if abs(a - b) > 1e-6 * max(1, abs(a)):
                                                    mism += 1
                                            elif (a is None) != (b is None) and not (a == '' or b == ''):
                                                if not (isinstance(a, str) and isinstance(b, str)):
                                                    mism += 1
                                        if mism or len(vis) != len(d['cache']):
                                            d['cacheMismatch'] = mism or f"len {len(vis)} vs {len(d['cache'])}"
                    fig['charts'].append(spec)
                for pic in anchor.iter('{%s}pic' % NS['xdr']):
                    blip = pic.find('.//a:blip', NS)
                    rid = blip.get(R + 'embed') if blip is not None else None
                    if rid and rid in drels:
                        media = drels[rid][1]
                        if media.endswith('image1.png'):
                            continue  # AEMO logo on every sheet
                        pic = {'file': posixpath.basename(media), 'from': pos(frm), 'to': pos(to)}
                        if frm is not None:
                            c0, r0 = pos(frm)
                            cap = grid.get(r0, c0 + 1)  # 0-based anchor row r0 -> 1-based row above it
                            if isinstance(cap, str) and cap.strip() and r0 > 2:
                                pic['caption'] = cap.strip()
                        fig['pictures'].append(pic)
        if fig['charts']:
            fig['kind'] = 'native'
        elif fig['pictures'] and len(grid.rows) > 22 and any(
                isinstance(v, (int, float)) for row in grid.rows[22:60] for v in row):
            fig['kind'] = 'picture-data'
        else:
            fig['kind'] = 'diagram'
        num = fig['number']
        if num and pages:
            n = int(re.match(r'\d+', num).group(0))
            pg = pages.get(fig['doc'], {}).get(str(n))
            if pg:
                fig['pdf'] = {'doc': fig['doc'], 'pages': pg,
                              'images': [f"{args.page_dir}/{fig['doc']}-p{p:03d}.png" for p in pg]}
        fig['grid'] = grid.rows
        fig['formats'] = grid.formats
        with open(os.path.join(args.out, fig['id'] + '.json'), 'w') as f:
            json.dump(fig, f, separators=(',', ':'), default=str)
        with open(os.path.join(args.out, fig['id'] + '.summary.txt'), 'w') as f:
            f.write(summarize(fig))
        index.append({k: fig.get(k) for k in ('id', 'sheet', 'kind', 'number', 'title', 'section', 'sectionName', 'doc', 'notes', 'pdf')}
                     | {'charts': [c['file'] for c in fig['charts']], 'pictures': [p['file'] for p in fig['pictures']],
                        'rows': len(grid.rows)})
        print(f"{fig['id']:12s} {fig['kind']:13s} charts={len(fig['charts'])} pics={len(fig['pictures'])}", file=sys.stderr)
    with open(os.path.join(args.out, 'index.json'), 'w') as f:
        json.dump(index, f, indent=1)


if __name__ == '__main__':
    main()
